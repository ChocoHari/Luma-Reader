"""Exercise packaged imports and native runtime without opening personal documents."""
import json
import sys
import tempfile
import os
import wave
from pathlib import Path


def check_reader_window(application,folder):
    """Exercise the real widget/editor tree without normal user state or hotkeys."""
    import app as reader_module
    import reader_core
    import studio
    from PySide6.QtCore import QEvent
    originals=[(module,'DATA',module.DATA) for module in (reader_module,reader_core,studio)]
    originals.append((reader_core,'VOICES',reader_core.VOICES))
    window=None
    try:
        for module in (reader_module,reader_core,studio):module.DATA=Path(folder)
        reader_core.VOICES=Path(folder)/'voices'
        window=reader_module.Reader(register_hotkey=False)
        window.new_note()
        window.editor.setPlainText('# Offline editor check\n\nA short original note.')
        window.refresh_markdown_workspace()
        window.ensurePolished()
        application.processEvents()
        result={'reader_window_rendered':not window.grab().isNull(),
                'markdown_editor':'Offline editor check' in window.markdown.toPlainText()}
        return result
    finally:
        if window is not None:
            window.editor.document().setModified(False)
            window.close();window.deleteLater()
            application.sendPostedEvents(None,QEvent.Type.DeferredDelete)
        for module,name,value in originals:setattr(module,name,value)


def run(output):
    report={'frozen':bool(getattr(sys,'frozen',False)),'checks':{},'errors':[]}
    try:
        import ssl
        report['runtime']={'python':sys.version.split()[0],'openssl':ssl.OPENSSL_VERSION}
        from runtime_privacy import prepare_inference_runtime
        prepare_inference_runtime()
        report['checks']['inference_telemetry_disabled']=os.environ.get('ORT_DISABLE_TELEMETRY')=='1'
        from PySide6.QtWidgets import QApplication
        from PySide6.QtMultimedia import QMediaPlayer,QAudioOutput,QAudioBufferOutput
        from PySide6.QtCore import QEventLoop,QTimer,QUrl
        import pymupdf
        from docx import Document
        from rapidocr_onnxruntime import RapidOCR
        from piper import PiperVoice
        from piper.phonemize_espeak import EspeakPhonemizer
        from runtime_paths import ASSET_ROOT,DATA_ROOT
        app=QApplication.instance() or QApplication([])
        player=QMediaPlayer()
        report['checks']['qt_platform']=app.platformName()
        report['checks']['assets']=all((ASSET_ROOT/p).is_file() for p in ('USER_GUIDE.md','assets/luma.ico','LICENSE'))
        source_manifest=json.loads((ASSET_ROOT/'SOURCE_MANIFEST.json').read_text())
        report['source_fingerprint']=source_manifest['source_fingerprint']
        report['checks']['source_identity']=len(report['source_fingerprint'])==64
        with tempfile.TemporaryDirectory(prefix='luma-diagnostic-') as folder:
            folder=Path(folder)
            path=folder/'sample.pdf'
            with pymupdf.open() as pdf:
                pdf.new_page().insert_text((50,80),'Luma standalone PDF check.')
                pdf.save(path)
            with pymupdf.open(path) as pdf:
                report['checks']['pdf']='Luma standalone' in pdf[0].get_text()
                report['checks']['pdf_render']=pdf[0].get_pixmap().width>0
            doc=Document();doc.add_paragraph('Luma document check.');doc.save(folder/'sample.docx')
            report['checks']['docx']=Document(folder/'sample.docx').paragraphs[0].text=='Luma document check.'
            report['checks'].update(check_reader_window(app,folder/'reader-state'))
        engine=RapidOCR()
        report['checks']['ocr_models_loaded']=engine is not None
        # Exercise OpenCV preprocessing and inference, not just model construction.
        from reader_core import OCR
        recognizer=OCR();recognizer.engine=engine
        with pymupdf.open() as sample:
            page=sample.new_page(width=500,height=180)
            page.insert_text((30,70),'Offline reading works.',fontsize=20)
            recognized=recognizer.recognize(page.get_pixmap(dpi=150,colorspace=pymupdf.csRGB,alpha=False))
        report['checks']['ocr_inference']='offline reading works' in recognized.casefold()
        report['checks']['piper_import']=PiperVoice is not None
        phonemizer=EspeakPhonemizer()
        report['checks']['piper_phonemization']=bool(phonemizer.phonemize('en-us','Luma standalone voice check.'))
        voice_path=os.environ.get('LUMA_TEST_VOICE')
        if voice_path:
            model=PiperVoice.load(voice_path)
            target=Path(output).with_suffix('.wav')
            with wave.open(str(target),'wb') as stream:
                model.synthesize_wav('Welcome to Luma. Your documents and your voice can stay on your computer.',stream)
            with wave.open(str(target),'rb') as stream:
                report['checks']['voice_synthesis']=stream.getnframes()>stream.getframerate()
                report['voice_seconds']=round(stream.getnframes()/stream.getframerate(),2)
                report['voice_sample_rate']=stream.getframerate()
            audio=QAudioOutput();audio.setMuted(True)
            player.setAudioOutput(audio)
            buffers=QAudioBufferOutput();player.setAudioBufferOutput(buffers)
            decoded=[];statuses=[]
            buffers.audioBufferReceived.connect(lambda buffer:decoded.append(buffer.frameCount()) if buffer.isValid() else None)
            loop=QEventLoop()
            def status_changed(status):
                statuses.append(status.name)
                if status in (QMediaPlayer.MediaStatus.EndOfMedia,QMediaPlayer.MediaStatus.InvalidMedia):loop.quit()
            player.mediaStatusChanged.connect(status_changed)
            player.errorOccurred.connect(lambda *args:loop.quit())
            timeout=QTimer();timeout.setSingleShot(True);timeout.timeout.connect(loop.quit)
            timeout.start(15000)
            player.setSource(QUrl.fromLocalFile(str(target.resolve())))
            player.play();loop.exec();timeout.stop()
            report['checks']['audio_decoded']=sum(decoded)>0
            report['checks']['playback_completed']='EndOfMedia' in statuses
            report['playback']={'muted':True,'decoded_frames':sum(decoded),'statuses':statuses,'error':player.errorString()}
            player.stop()
        DATA_ROOT.mkdir(parents=True,exist_ok=True)
        report['checks']['user_storage']=DATA_ROOT.is_dir()
    except Exception as error:
        import traceback
        report['errors'].append(traceback.format_exc())
    report['passed']=not report['errors'] and all(report['checks'].values())
    Path(output).write_text(json.dumps(report,indent=2),encoding='utf-8')
    return 0 if report['passed'] else 1
