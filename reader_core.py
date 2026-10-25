"""Local document extraction, voice discovery, and persistent reader state."""
from __future__ import annotations

import hashlib
import json
import os
import re
import tempfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Callable

from runtime_paths import ASSET_ROOT as ROOT, DATA_ROOT as DATA, VOICE_ROOT as VOICES
from runtime_privacy import prepare_inference_runtime


def atomic_json(path: Path, value) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, name = tempfile.mkstemp(dir=path.parent, suffix=".tmp")
    try:
        with os.fdopen(fd, "w", encoding="utf-8") as stream:
            json.dump(value, stream, ensure_ascii=False, indent=2)
        os.replace(name, path)
    finally:
        Path(name).unlink(missing_ok=True)


def read_json(path: Path, default):
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return default


@dataclass
class Voice:
    path: Path
    name: str
    speakers: dict[str, int] = field(default_factory=dict)


def discover_voices(folder: Path | None = None) -> tuple[list[Voice], list[str]]:
    voices, issues = [], []
    folders=[Path(folder)] if folder is not None else [VOICES,ROOT/'voices',ROOT/'Piper Voices']
    if folder is not None:Path(folder).mkdir(parents=True,exist_ok=True)
    paths=sorted({path.resolve() for location in folders if location.is_dir() for path in location.rglob('*.onnx')})
    for path in paths:
        try:
            config = json.loads(Path(str(path) + ".json").read_text(encoding="utf-8"))
            if not config.get("audio", {}).get("sample_rate") or not config.get("phoneme_id_map"):
                raise ValueError("configuration is missing audio or phoneme information")
            count = int(config.get("num_speakers", 1))
            speakers = config.get("speaker_id_map") or {f"Speaker {i + 1}": i for i in range(count)}
            voices.append(Voice(path, path.stem.replace("en_GB-", "").replace("_", " "), speakers))
        except (OSError, ValueError, TypeError) as error:
            issues.append(f"{path.name}: {error}")
    return voices, issues


@dataclass
class Page:
    text: str
    method: str = "Text"
    regions: list = field(default_factory=list)


@dataclass
class Document:
    path: Path | None
    pages: list[Page]
    title: str
    headings: list[tuple[str, int]] = field(default_factory=list)
    rich_html: str | None = None

    @property
    def text(self) -> str:
        return "\n\n".join(page.text for page in self.pages)

    def page_start(self, page: int) -> int:
        return sum(len(p.text) + 2 for p in self.pages[:page])


class Cancelled(Exception):
    pass


def clean_pdf_text(text: str) -> str:
    text = text.replace("\x00", "").replace("\u00ad", "")
    text = re.sub(r"(?<=[a-zA-Z])-\n(?=[a-z])", "", text)
    paragraphs = re.split(r"\n\s*\n", text)
    return "\n\n".join(re.sub(r"[ \t]*\n[ \t]*", " ", p).strip() for p in paragraphs).strip()


def pdf_key(path: Path) -> str:
    stat = path.stat()
    return hashlib.sha256(f"{path.resolve()}|{stat.st_size}|{stat.st_mtime_ns}|ocr-v1".encode()).hexdigest()


class OCR:
    """Packaged ONNX models; no service or first-run model download."""
    def __init__(self):
        self.engine = None
        self.boxes = []

    def recognize(self, pixmap) -> str:
        import numpy as np
        if self.engine is None:
            from ocr_backend import create_engine
            self.engine = create_engine()
        rgb = np.frombuffer(pixmap.samples, dtype=np.uint8).reshape(pixmap.height, pixmap.width, pixmap.n)
        # RapidOCR expects BGR for a numpy input.
        image = rgb[:, :, :3][:, :, ::-1].copy()
        result, _ = self.engine(image)
        self.boxes = []
        if not result:
            return ""
        # Detection order retains lines; it cannot infer every multi-column layout.
        accepted = [item for item in result if float(item[2]) >= 0.45]
        self.boxes = [(item[0], str(item[1])) for item in accepted]
        return "\n".join(str(item[1]) for item in accepted)


def ocr_regions(ocr, text, pix, rect):
    regions, position = [], 0
    for box, label in getattr(ocr, "boxes", []):
        start = text.find(label, position)
        if start < 0:
            continue
        xs, ys = [p[0] for p in box], [p[1] for p in box]
        regions.append({"start": start, "end": start + len(label), "rect":
            [min(xs)*rect.width/pix.width, min(ys)*rect.height/pix.height,
             max(xs)*rect.width/pix.width, max(ys)*rect.height/pix.height]})
        position = start + len(label)
    return regions


def load_document(path: Path, ocr_mode="auto", progress: Callable | None = None,
                  cancelled: Callable = lambda: False, ocr: OCR | None = None) -> Document:
    def report(current, total, message):
        if progress:
            progress(current, total, message)
        if cancelled():
            raise Cancelled()

    suffix = path.suffix.lower()
    if suffix == ".pdf":
        import pymupdf as fitz
        pages = []
        ocr = ocr or OCR()
        cache_dir = DATA / "ocr" / pdf_key(path)
        with fitz.open(path) as pdf:
            if pdf.needs_pass:
                raise ValueError("This PDF is password protected. Open an unlocked copy.")
            toc = [(title, max(0, page - 1)) for _, title, page, *_ in pdf.get_toc()]
            for i, page in enumerate(pdf):
                report(i, len(pdf), f"Reading page {i + 1} of {len(pdf)}")
                cached = read_json(cache_dir / f"{i}.json", None)
                from pdf_reading_order import native_page_content
                native_text, _, reading_order = native_page_content(page, include_words=False)
                text = clean_pdf_text(native_text)
                needs_ocr = ocr_mode == "force" or (ocr_mode == "auto" and len(re.sub(r"\s", "", text)) < 30)
                method = "Text · two columns" if reading_order == 'Two-column prose' else "Text"
                regions = []
                legacy_cache = cached and cached.get("text","").strip() and "regions" not in cached
                if cached and ocr_mode != "off" and not legacy_cache:
                    text, method = cached["text"], "OCR (cached)"
                    regions = cached.get("regions", [])
                elif needs_ocr or (legacy_cache and ocr_mode != "off"):
                    report(i, len(pdf), f"Recognizing page {i + 1} of {len(pdf)} · offline OCR")
                    pix = page.get_pixmap(dpi=180, colorspace=fitz.csRGB, alpha=False)
                    text = ocr.recognize(pix).strip()
                    method = "OCR"
                    regions = ocr_regions(ocr, text, pix, page.rect)
                    atomic_json(cache_dir / f"{i}.json", {"text": text, "regions": regions})
                if 'OCR' in method:
                    from pdf_reading_order import ordered_ocr_regions
                    text, regions, columns = ordered_ocr_regions(text, regions, page.rect.width)
                    if columns:
                        method += ' · two columns'
                        if not cached or cached.get('text') != text or cached.get('regions') != regions:
                            atomic_json(cache_dir / f"{i}.json", {"text": text, "regions": regions})
                pages.append(Page(text, method, regions))
            report(len(pdf), len(pdf), "Ready")
        return Document(path, pages, path.name, toc)
    if suffix == ".docx":
        from docx import Document as WordDocument
        from docx.table import Table
        from docx.text.paragraph import Paragraph
        word = WordDocument(path)
        parts, headings = [], []
        for block in word.iter_inner_content():
            if cancelled():
                raise Cancelled()
            if isinstance(block, Paragraph):
                if block.style and block.style.name.startswith("Heading"):
                    headings.append((block.text, 0))
                parts.append(block.text)
            elif isinstance(block, Table):
                parts.extend(" | ".join(cell.text for cell in row.cells) for row in block.rows)
        return Document(path, [Page("\n\n".join(parts))], path.name, headings)
    if suffix in {".txt", ".md", ".markdown"}:
        raw = path.read_bytes()
        if raw.startswith((b"\xff\xfe", b"\xfe\xff")):
            text = raw.decode("utf-16")
        else:
            try:
                text = raw.decode("utf-8-sig")
            except UnicodeDecodeError:
                text = raw.decode("cp1252")
        return Document(path, [Page(text.replace("\r\n", "\n").replace("\r", "\n"))], path.name)
    if suffix in {".html", ".htm"}:
        from html.parser import HTMLParser
        class NoteText(HTMLParser):
            def __init__(self):
                super().__init__()
                self.parts, self.hidden = [], 0
            def handle_starttag(self, tag, attrs):
                if tag in {"script", "style", "head"}:
                    self.hidden += 1
                if tag in {"p", "div", "br", "li", "h1", "h2", "h3"} and not self.hidden:
                    self.parts.append("\n")
            def handle_endtag(self, tag):
                if tag in {"script", "style", "head"}:
                    self.hidden = max(0, self.hidden-1)
            def handle_data(self, data):
                if not self.hidden:
                    self.parts.append(data)
        raw = path.read_text(encoding="utf-8-sig")
        parser = NoteText()
        parser.feed(raw)
        return Document(path, [Page("".join(parser.parts).strip())], path.name, rich_html=raw)
    raise ValueError("Open a PDF, DOCX, TXT, or Markdown file.")


def ocr_pdf_page(path: Path, index: int, ocr: OCR | None = None) -> Page:
    import pymupdf as fitz
    with fitz.open(path) as pdf:
        pix = pdf[index].get_pixmap(dpi=200, colorspace=fitz.csRGB, alpha=False)
        engine = ocr or OCR()
        text = engine.recognize(pix).strip()
        regions = ocr_regions(engine, text, pix, pdf[index].rect)
        from pdf_reading_order import ordered_ocr_regions
        text, regions, columns = ordered_ocr_regions(text, regions, pdf[index].rect.width)
    atomic_json(DATA / "ocr" / pdf_key(path) / f"{index}.json", {"text": text, "regions": regions})
    return Page(text, "OCR · two columns" if columns else "OCR", regions)


@dataclass
class Sentence:
    start: int
    end: int
    text: str
    spoken: str | None = None
    pause_after: float = 0.0


def sentences(text: str, start=0, end: int | None = None) -> list[Sentence]:
    end = len(text) if end is None else end
    output = []
    # Keep source offsets intact. Avoid splitting initials, decimals, and common titles.
    pattern = r"[.!?]+[\"'\u201d\u2019)]*(?=\s|$)|\n\s*\n"
    position = start
    for match in re.finditer(pattern, text[start:end]):
        boundary = start + match.end()
        segment = text[position:boundary]
        if re.search(r"\b(?:Mr|Mrs|Ms|Dr|Prof|Sr|Jr|St|vs|etc|Fig|Figs|No|e\.g|i\.e)\.$|\b[A-Z]\.$", segment.strip(), re.I):
            continue
        _append_chunks(output, text, position, boundary)
        position = boundary
    _append_chunks(output, text, position, end)
    return output


def _append_chunks(output, text, start, end):
    while start < end:
        while start < end and text[start].isspace():
            start += 1
        if start >= end:
            return
        stop = min(end, start + 600)
        if stop < end:
            space = text.rfind(" ", start + 200, stop)
            if space > start:
                stop = space + 1
        trimmed = stop
        while trimmed > start and text[trimmed - 1].isspace():
            trimmed -= 1
        if trimmed > start:
            output.append(Sentence(start, trimmed, text[start:trimmed]))
        start = stop


def qt_position(text: str, python_position: int) -> int:
    return len(text[:python_position].encode("utf-16-le")) // 2


def python_position(text: str, qt_offset: int) -> int:
    return len(text.encode("utf-16-le")[:qt_offset * 2].decode("utf-16-le", errors="ignore"))


class State:
    def __init__(self):
        self.path = DATA / "settings.json"
        self.value = read_json(self.path, {})

    def save(self):
        atomic_json(self.path, self.value)

    def remember(self, path: Path, position: int):
        key = str(path.resolve())
        self.value.setdefault("positions", {})[key] = position
        recent = self.value.setdefault("recent", [])
        self.value["recent"] = ([key] + [p for p in recent if p != key])[:15]
        self.save()
