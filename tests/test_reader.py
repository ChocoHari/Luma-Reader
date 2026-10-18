"""Functional checks against actual engines and temporary document fixtures."""
import os
os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

import json
import socket
import time
import wave
from pathlib import Path

import pytest

import reader_core as core


@pytest.fixture(autouse=True)
def private_data(tmp_path, monkeypatch):
    monkeypatch.setattr(core, "DATA", tmp_path / "state")


def test_voice_pairs():
    voices, issues = core.discover_voices()
    assert voices, 'Install the reviewed build voice with tools/fetch_build_voice.py before running engine tests'
    assert not issues
    assert all(voice.speakers for voice in voices)


def test_unicode_sentence_offsets_and_long_text():
    text = "Dr. Ada reads 😀. Next sentence!\n\n" + "Long passage " * 180
    parts = core.sentences(text)
    assert parts[0].text == "Dr. Ada reads 😀."
    assert parts[1].text == "Next sentence!"
    assert all(part.text == text[part.start:part.end] and len(part.text) <= 600 for part in parts)
    assert core.python_position(text, core.qt_position(text, 16)) == 16
    assert "".join(part.text.replace(" ", "").replace("\n", "") for part in parts) == text.replace(" ", "").replace("\n", "")


def test_text_encodings_and_state(tmp_path):
    path = tmp_path / "notes.md"
    path.write_text("# My notes\n\nCafé and 😀", encoding="utf-8-sig")
    assert "Café" in core.load_document(path).text
    path.write_text("Unicode notes", encoding="utf-16")
    assert core.load_document(path).text == "Unicode notes"
    state = core.State()
    state.remember(path, 5)
    assert core.State().value["positions"][str(path.resolve())] == 5


def test_docx_paragraph_table_order(tmp_path):
    from docx import Document
    word = Document()
    word.add_heading("Overview", 1)
    word.add_paragraph("Before the table.")
    table = word.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Left cell"
    table.cell(0, 1).text = "Right cell"
    word.add_paragraph("After the table.")
    path = tmp_path / "test.docx"
    word.save(path)
    document = core.load_document(path)
    assert document.text.index("Before") < document.text.index("Left cell") < document.text.index("After")
    assert document.headings == [("Overview", 0)]


def test_pdf_native_text_and_cancellation(tmp_path):
    import pymupdf as fitz
    path = tmp_path / "text.pdf"
    with fitz.open() as pdf:
        pdf.new_page().insert_text((50, 80), "A document with selectable text and a readable first sentence.")
        pdf.new_page().insert_text((50, 80), "A second readable page with its own selectable text.")
        pdf.save(path)
    document = core.load_document(path, "off")
    assert len(document.pages) == 2
    assert "first sentence" in document.pages[0].text
    assert document.text[document.page_start(1):].startswith("A second")
    with pytest.raises(core.Cancelled):
        core.load_document(path, cancelled=lambda: True)


def forbid_network(monkeypatch):
    def denied(*args, **kwargs):
        raise AssertionError("An offline engine attempted to connect to a network.")
    monkeypatch.setattr(socket.socket, "connect", denied)
    monkeypatch.setattr(socket, "create_connection", denied)


def test_actual_offline_ocr_and_cache(tmp_path, monkeypatch):
    import pymupdf as fitz
    from PIL import Image, ImageDraw, ImageFont
    image = Image.new("RGB", (1400, 700), "white")
    drawing = ImageDraw.Draw(image)
    font = ImageFont.truetype("C:/Windows/Fonts/arial.ttf", 46)
    drawing.text((70, 90), "Offline reading works.", font=font, fill="black")
    drawing.text((70, 190), "Figure 1. Water flows through a pipe.", font=font, fill="black")
    drawing.rectangle((100, 350, 1000, 430), outline="black", width=4)
    drawing.text((450, 470), "Water pipe", font=font, fill="black")
    import io
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    path = tmp_path / "scanned.pdf"
    with fitz.open() as pdf:
        page = pdf.new_page(width=700, height=350)
        page.insert_image(page.rect, stream=buffer.getvalue())
        pdf.save(path)
    forbid_network(monkeypatch)
    document = core.load_document(path)
    assert "offline reading works" in document.text.lower()
    assert "water" in document.text.lower()
    assert document.pages[0].method == "OCR"
    class NoRecognition:
        def recognize(self, _):
            raise AssertionError("Cached OCR should not run inference again")
    cached = core.load_document(path, ocr=NoRecognition())
    assert cached.text == document.text
    assert "cached" in cached.pages[0].method
    assert core.load_document(path, "off").text == ""


def test_actual_offline_two_column_scan_preserves_line_regions(tmp_path,monkeypatch):
    import pymupdf
    path=tmp_path/'scanned-columns.pdf'
    with pymupdf.open() as original:
        page=original.new_page(width=600,height=400)
        page.insert_text((40,35),'Original scanned reading example',fontsize=18)
        for n in range(1,6):
            for x,label in ((40,'Left'),(330,'Right')):
                page.insert_text((x,70+n*25),f'{label} line {n} has several original words.',fontsize=11)
        image=page.get_pixmap(dpi=200,alpha=False).tobytes('png')
    with pymupdf.open() as scan:
        page=scan.new_page(width=600,height=400)
        page.insert_image(page.rect,stream=image)
        scan.save(path)
    forbid_network(monkeypatch)
    document=core.load_document(path)
    page=document.pages[0]
    assert page.method=='OCR · two columns'
    assert document.text.index('Left line 5') < document.text.index('Right line 1')
    assert len(page.regions)==11
    for region in page.regions:
        label=page.text[region['start']:region['end']]
        if label.startswith('Left'):
            assert region['rect'][2]<300
        elif label.startswith('Right'):
            assert region['rect'][0]>300


@pytest.mark.parametrize("voice", core.discover_voices()[0], ids=lambda voice:voice.name)
def test_actual_piper_voices_offline(tmp_path, monkeypatch, voice):
    from speech import VoiceEngine
    import numpy as np
    forbid_network(monkeypatch)
    target = tmp_path / "voice.wav"
    VoiceEngine().write("This voice is ready for offline reading.", target, voice)
    with wave.open(str(target), "rb") as wav:
        assert wav.getnchannels() == 1
        assert wav.getsampwidth() == 2
        assert wav.getnframes() / wav.getframerate() > 1
        samples = np.frombuffer(wav.readframes(wav.getnframes()), dtype=np.int16)
        assert np.max(np.abs(samples.astype(np.int32))) > 500


def test_export_cancel_preserves_destination(tmp_path):
    from speech import export_audio
    target = tmp_path / "existing.wav"
    target.write_bytes(b"existing file")
    voice = core.discover_voices()[0][0]
    result = export_audio(core.sentences("First sentence. Second sentence."), voice, target, 1, 0, lambda *a: None, lambda: True)
    assert result is False
    assert target.read_bytes() == b"existing file"


def test_export_real_wav(tmp_path):
    from speech import export_audio
    target = tmp_path / "export.wav"
    voice = core.discover_voices()[0][0]
    events = []
    assert export_audio(core.sentences("First sentence. Second sentence."), voice, target, 1.2, 0, lambda *a: events.append(a), lambda: False)
    assert len(events) == 2
    with wave.open(str(target), "rb") as wav:
        assert wav.getnframes() > wav.getframerate()


@pytest.fixture(scope="module")
def qt_app():
    from PySide6.QtWidgets import QApplication
    app = QApplication.instance() or QApplication([])
    yield app


def test_gui_note_save_unicode_selection_and_restore(tmp_path, monkeypatch, qt_app):
    from app import Reader
    from PySide6.QtGui import QTextCursor
    from PySide6.QtWidgets import QFileDialog
    window = Reader(register_hotkey=False)
    window.new_plain()
    text = "A note with 😀. A second sentence."
    window.editor.insertPlainText(text)
    assert window.editor.document().isModified()
    path = tmp_path / "new-note.md"
    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *a, **k: (str(path), ""))
    assert window.save()
    assert path.read_text(encoding="utf-8") == text
    assert not window.editor.document().isModified()
    window.set_cursor(text.index("A second"))
    window.save_progress()
    window.apply_document(core.load_document(path))
    assert core.python_position(text, window.editor.textCursor().position()) == text.index("A second")
    captured = []
    monkeypatch.setattr(window.speech, "start", lambda *a: captured.append(a))
    window.play_pause()
    assert captured[-1][2] == 1
    cursor = window.editor.textCursor()
    cursor.setPosition(core.qt_position(text, text.index("A second")))
    cursor.movePosition(QTextCursor.MoveOperation.End, QTextCursor.MoveMode.KeepAnchor)
    window.editor.setTextCursor(cursor)
    window.read_selection()
    assert window.parts[0].text == "A second sentence."
    window.highlight(0)
    assert window.editor.extraSelections()[0].cursor.selectedText() == "A second sentence."
    window.close()


def test_gui_stop_discards_test_scope_and_pdf_preview(tmp_path, qt_app):
    import pymupdf as fitz
    from app import Reader
    window = Reader(register_hotkey=False)
    window.testing_voice = True
    window.parts = core.sentences("Test voice text.")
    window.stop()
    assert not window.parts
    path = tmp_path / "preview.pdf"
    with fitz.open() as pdf:
        pdf.new_page().insert_text((50, 80), "Page preview test, selectable text and normal reading.")
        pdf.save(path)
    window.apply_document(core.load_document(path, "off"))
    window.show()
    qt_app.processEvents()
    window.render_page()
    assert window.preview_label.pixmap() and not window.preview_label.pixmap().isNull()
    assert window.editor.isReadOnly()
    window.close()


def test_speech_queue_stop_cancels_producer(tmp_path, monkeypatch, qt_app):
    from speech import Speech
    controller = Speech()
    entered = threading_event = __import__("threading").Event()
    def fake_write(text, path, *args):
        threading_event.set()
        with wave.open(str(path), "wb") as wav:
            wav.setnchannels(1)
            wav.setsampwidth(2)
            wav.setframerate(22050)
            wav.writeframes(b"\0\0" * 2205)
    monkeypatch.setattr(controller.engine, "write", fake_write)
    controller.start(core.sentences("Sentence one. Sentence two. Sentence three. Sentence four. Sentence five."), core.discover_voices()[0][0])
    assert entered.wait(5)
    folder = controller.folder
    controller.stop()
    controller.thread.join(timeout=5)
    controller._clean_retired()
    assert not controller.thread.is_alive()
    assert not controller.active
    assert not folder.exists()
    controller.timer.stop()


def test_actual_audio_pause_resume_and_completion(qt_app):
    from speech import Speech
    from PySide6.QtMultimedia import QMediaPlayer
    controller = Speech()
    controller.output.setVolume(0)
    starts, errors, completed = [], [], []
    controller.sentence_started.connect(starts.append)
    controller.failed.connect(errors.append)
    controller.finished.connect(lambda: completed.append(True))
    controller.start(core.sentences("Hello reader. Goodbye reader."), core.discover_voices()[0][0])
    deadline = time.monotonic() + 30
    while not starts and not errors and time.monotonic() < deadline:
        qt_app.processEvents()
        time.sleep(0.01)
    assert starts and not errors
    controller.toggle_pause()
    assert controller.paused
    before = len(starts)
    paused_until = time.monotonic() + 0.15
    while time.monotonic() < paused_until:
        qt_app.processEvents()
        time.sleep(0.01)
    assert len(starts) == before
    controller.toggle_pause()
    assert not controller.paused
    while not completed and not errors and time.monotonic() < deadline:
        qt_app.processEvents()
        time.sleep(0.01)
    controller.stop()
    controller.timer.stop()
    assert not errors
    assert completed and starts == [0, 1]

